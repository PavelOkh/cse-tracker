import io
import re
import time
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import PatternFill, Font
import pandas as pd
import requests
import streamlit as st

st.set_page_config(page_title="CSE Трекинг", page_icon="📦", layout="wide")

def process_single_number(number):
    url = f'https://lk.cse.ru/api/new-track/{number}'
    headers = {
        'Accept': 'application/json, text/plain, */*',
        'Origin': 'https://www.cse.ru',
        'Referer': 'https://www.cse.ru/',
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/150.0.0.0 YaBrowser/26.8.0.0 Safari/537.36'
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
        data = response.json()
        
        found_items = data.get('found', [])
        if not found_items:
            return {
                "Накладная": number, "Статус": "Данные не найдены", "Дата доставки": "",
                "Дата посл. статуса": "", "Изначальный номер": "", "Новый номер": "", "Статус нового номера": "",
                "Дата доставки нового": "", "Дата статуса нового": "", "Тип подсветки": "желтый"
            }
            
        track_data = found_items[0]
        
        state = track_data.get('State', 'Статус не указан')
        info = track_data.get('Info', '')
        
        current_status = f"{state}. {info}" if info else state
        delivery_date = ""
        last_status_date = ""
        original_waybill_number = ""
        new_waybill_number = ""
        new_state = ""
        new_delivery_date = ""
        new_last_status_date = ""
        
        highlight_type = "желтый"

        history_dict = track_data.get('History', {})
        waybill_info = history_dict.get('waybill_info', [])
        order_info = history_dict.get('order_info', [])
        
        # 1. Поиск дочернего документа (если текущий номер — родитель)
        for event in waybill_info:
            doc = event.get('Document')
            if doc and isinstance(doc, dict):
                doc_number = doc.get('Number', '')
                doc_state = doc.get('State', '')
                if doc_number and doc_number != number:
                    new_waybill_number = doc_number
                    new_state = doc_state
                    
                    doc_history = doc.get('History', [])
                    if doc_history:
                        new_last_status_date = doc_history[0].get('EventDate', '')
                        
                    for sub_event in doc_history:
                        if "Доставка завершена" in sub_event.get('EventName', ''):
                            new_delivery_date = sub_event.get('EventDate', '')
                            break
                    break

            event_name = event.get('EventName', '')
            event_name_en = event.get('EventNameEn', '')
            event_date = event.get('EventDate', '')
            
            if "Доставка завершена" in event_name or event_name_en == "Delivery completed":
                current_status = "Доставка завершена"
                delivery_date = event_date
                last_status_date = ""
                highlight_type = "нет"
                break

        if not new_waybill_number and ("возврат" in state.lower() or "возвращается" in current_status.lower() or "смена" in current_status.lower()):
            parts = current_status.split()
            for p in parts:
                cleaned_p = p.strip(".,;")
                if "497-" in cleaned_p and cleaned_p != number:
                    new_waybill_number = cleaned_p
                    break

        # 2. Если дочерний номер не найден — проверяем, является ли номер дочерним (ищем родителя в истории)
        if not new_waybill_number:
            all_events = waybill_info + order_info
            for event in all_events:
                ev_info = event.get("EventInfo", "")
                matches = re.findall(r'(497-[\d\-A-Z]+)', ev_info)
                for match in matches:
                    cleaned = match.strip(".,;")
                    if cleaned != number:
                        original_waybill_number = cleaned
                        break
                if original_waybill_number:
                    break

        if waybill_info and not delivery_date:
            last_status_date = waybill_info[0].get('EventDate', '')

        if "возврат" in state.lower() or "возвращается" in current_status.lower():
            highlight_type = "красный"
        elif new_waybill_number or "добавочная" in current_status.lower() or "смена" in current_status.lower():
            highlight_type = "фиолетовый"

        if current_status == "Доставка завершена":
            highlight_type = "нет"

        return {
            "Накладная": number,
            "Статус": current_status,
            "Дата доставки": delivery_date,
            "Дата посл. статуса": last_status_date if current_status != "Доставка завершена" else "",
            "Изначальный номер": original_waybill_number,
            "Новый номер": new_waybill_number,
            "Статус нового номера": new_state,
            "Дата доставки нового": new_delivery_date,
            "Дата статуса нового": new_last_status_date if (new_state and new_state != "Доставка завершена") else "",
            "Тип подсветки": highlight_type
        }

    except Exception:
        return {
            "Накладная": number, "Статус": "Ошибка обработки", "Дата доставки": "",
            "Дата посл. статуса": "", "Изначальный номер": "", "Новый номер": "", "Статус нового номера": "",
            "Дата доставки нового": "", "Дата статуса нового": "", "Тип подсветки": "ошибка"
        }

def fetch_with_smart_retries(unique_numbers, progress_callback):
    """Многопроходная обработка с отложенными повторными запросами (до 5 попыток)"""
    results_dict = {}
    current_batch = list(unique_numbers)
    total_unique = len(unique_numbers)
    
    for attempt in range(1, 6):
        if not current_batch:
            break
            
        if attempt > 1:
            # Увеличенный интервал ожидания для отложенного списка (2с, 5с, 10с, 15с)
            wait_time = [0, 2, 5, 10, 15][attempt - 1]
            progress_callback(None, f"Попытка {attempt}/5: повторный опрос {len(current_batch)} накладных (ожидание {wait_time}с)...")
            time.sleep(wait_time)
        
        failed_batch = []
        completed_in_pass = 0
        
        with ThreadPoolExecutor(max_workers=10) as executor:
            future_to_num = {executor.submit(process_single_number, num): num for num in current_batch}
            for future in as_completed(future_to_num):
                num = future_to_num[future]
                completed_in_pass += 1
                try:
                    res = future.result()
                    if res["Статус"] == "Ошибка обработки" and attempt < 5:
                        failed_batch.append(num)
                    results_dict[num] = res
                except Exception:
                    if attempt < 5:
                        failed_batch.append(num)
                    results_dict[num] = {
                        "Накладная": num, "Статус": "Ошибка обработки", "Дата доставки": "",
                        "Дата посл. статуса": "", "Изначальный номер": "", "Новый номер": "", "Статус нового номера": "",
                        "Дата доставки нового": "", "Дата статуса нового": "", "Тип подсветки": "ошибка"
                    }
                
                # Общий прогресс текущего прохода
                processed_count = len(unique_numbers) - len(current_batch) + completed_in_pass
                progress_callback(processed_count / total_unique, f"Обработка (попытка {attempt}/5): {processed_count}/{total_unique}")
                
        current_batch = failed_batch
        
    return results_dict

def generate_colored_excel(df_data):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_data.to_excel(writer, index=False, sheet_name='Результат')
    
    output.seek(0)
    wb = openpyxl.load_workbook(output)
    ws = wb.active
    
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    purple_fill = PatternFill(start_color="E1D5E7", end_color="E1D5E7", fill_type="solid")
    red_fill = PatternFill(start_color="F8CECC", end_color="F8CECC", fill_type="solid")
    grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    red_font = Font(color="9C0006", bold=True)
    
    headers = [cell.value for cell in ws[1]]
    type_col_idx = headers.index("Тип подсветки") + 1
    
    for row_idx in range(2, ws.max_row + 1):
        highlight_type = ws.cell(row=row_idx, column=type_col_idx).value
        target_fill = None
        target_font = None
        
        if highlight_type == "желтый":
            target_fill = yellow_fill
        elif highlight_type == "фиолетовый":
            target_fill = purple_fill
        elif highlight_type == "красный":
            target_fill = red_fill
        elif highlight_type == "ошибка":
            target_fill = grey_fill
            target_font = red_font
            
        if target_fill or target_font:
            for col_idx in range(1, len(headers)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if target_fill:
                    cell.fill = target_fill
                if target_font:
                    cell.font = target_font
                
    ws.delete_cols(type_col_idx)
    
    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()

def color_rows(row):
    h_type = row["Тип подсветки"]
    if h_type == "желтый":
        return ['background-color: #fff2cc'] * len(row)
    elif h_type == "фиолетовый":
        return ['background-color: #e1d5e7'] * len(row)
    elif h_type == "красный":
        return ['background-color: #f8cecc'] * len(row)
    elif h_type == "ошибка":
        return ['background-color: #d9d9d9; color: #9c0006; font-weight: bold'] * len(row)
    return [''] * len(row)

def main():
    st.title("📦 Массовая проверка накладных CSE")
    st.write("Инструмент автоматического отслеживания отправлений с умными повторными попытками.")

    with st.expander("🎨 Справка по цветовой индикации"):
        st.markdown("""
        * 🟨 **Желтый фон** — доставка в процессе / не завершена.
        * 🟪 **Фиолетовый фон** — смена номера накладной (досыл / добавочная).
        * 🟥 **Красный фон** — оформлен возврат отправителю.
        * ⬜ **Серо-красный** — ошибка обработки (система автоматически повторяла попытку до 5 раз).
        * *Без заливки* — доставка успешно завершена.
        """)

    input_method = st.radio("Выберите способ ввода данных:", ["📁 Загрузить файл (Excel / CSV)", "📋 Вставить списком из буфера обмена"])
    
    tracking_list = []
    
    if input_method == "📁 Загрузить файл (Excel / CSV)":
        uploaded_file = st.file_uploader("Перетащите файл сюда или нажмите Browse", type=["xlsx", "xls", "csv"])
        if uploaded_file is not None:
            try:
                if uploaded_file.name.endswith('.csv'):
                    try:
                        df_input = pd.read_csv(uploaded_file, encoding='utf-8')
                    except Exception:
                        df_input = pd.read_csv(uploaded_file, encoding='cp1251')
                else:
                    df_input = pd.read_excel(uploaded_file)
                
                columns = df_input.columns.tolist()
                selected_column = st.selectbox("Выберите столбец с номерами накладных:", columns, index=0)
                tracking_list = df_input[selected_column].tolist()
            except Exception as e:
                st.error(f"Ошибка при чтении файла: {e}")
    else:
        raw_text = st.text_area("Вставьте список накладных (разделители: перенос строки, табуляция, запятая, точка с запятой):", height=150)
        if raw_text:
            items = re.split(r'[\r\n,\t;]+', raw_text)
            tracking_list = [item.strip() for item in items if item.strip()]
            if len(tracking_list) > 1000:
                st.warning("⚠️ Внимание: вставлено более 1000 номеров. Будут обработаны первые 1000.")
                tracking_list = tracking_list[:1000]

    if tracking_list:
        original_list = [str(num).strip() if pd.notna(num) else "" for num in tracking_list]
        unique_numbers = list(dict.fromkeys([num for num in original_list if num]))
        
        st.info(f"Всего строк для обработки: **{len(original_list)}** | Уникальных номеров для API: **{len(unique_numbers)}**")
        
        if st.button("🚀 Начать проверку", type="primary"):
            start_time = time.time()
            my_bar = st.progress(0, text="Идет опрос API CSE...")
            
            def update_progress(val, text):
                if val is not None:
                    my_bar.progress(val, text=text)
                else:
                    my_bar.progress(0.0, text=text)
            
            results_dict = fetch_with_smart_retries(unique_numbers, update_progress)
            my_bar.empty()
            
            final_results = []
            for num in original_list:
                if num == "":
                    final_results.append({
                        "Накладная": "", "Статус": "Пустая строка", "Дата доставки": "",
                        "Дата посл. статуса": "", "Изначальный номер": "", "Новый номер": "", "Статус нового номера": "",
                        "Дата доставки нового": "", "Дата статуса нового": "", "Тип подсветки": "нет"
                    })
                else:
                    final_results.append(results_dict.get(num, {
                        "Накладная": num, "Статус": "Ошибка обработки", "Дата доставки": "",
                        "Дата посл. статуса": "", "Изначальный номер": "", "Новый номер": "", "Статус нового номера": "",
                        "Дата доставки нового": "", "Дата статуса нового": "", "Тип подсветки": "ошибка"
                    }))
            
            elapsed_time = round(time.time() - start_time, 1)
            st.success(f"✅ Проверка завершена за {elapsed_time} сек.")
            
            df_output = pd.DataFrame(final_results)
            
            # Сводные метрики (теперь 6 колонок)
            total_count = len(final_results)
            delivered_count = sum(1 for r in final_results if r["Тип подсветки"] == "нет" and r["Статус"] == "Доставка завершена")
            yellow_count = sum(1 for r in final_results if r["Тип подсветки"] == "желтый")
            purple_count = sum(1 for r in final_results if r["Тип подсветки"] == "фиолетовый")
            red_count = sum(1 for r in final_results if r["Тип подсветки"] == "красный")
            error_count = sum(1 for r in final_results if r["Тип подсветки"] == "ошибка")
            
            m1, m2, m3, m4, m5, m6 = st.columns(6)
            m1.metric("📦 Всего", total_count)
            m2.metric("✅ Доставлено", delivered_count)
            m3.metric("🟨 В пути", yellow_count)
            m4.metric("🟪 Смена номера", purple_count)
            m5.metric("🟥 Возвраты", red_count)
            m6.metric("⬜ Ошибки", error_count)
            
            # Фильтрация отображения
            display_filter = st.selectbox("Фильтр отображения в таблице ниже:", [
                "Все строки", 
                "Только в пути (желтые)", 
                "Смена номера (фиолетовые)", 
                "Возвраты (красные)", 
                "Ошибки обработки",
                "Доставленные"
            ])
            
            df_filtered = df_output.copy()
            if display_filter == "Только в пути (желтые)":
                df_filtered = df_output[df_output["Тип подсветки"] == "желтый"]
            elif display_filter == "Смена номера (фиолетовые)":
                df_filtered = df_output[df_output["Тип подсветки"] == "фиолетовый"]
            elif display_filter == "Возвраты (красные)":
                df_filtered = df_output[df_output["Тип подсветки"] == "красный"]
            elif display_filter == "Ошибки обработки":
                df_filtered = df_output[df_output["Тип подсветки"] == "ошибка"]
            elif display_filter == "Доставленные":
                df_filtered = df_output[df_output["Статус"] == "Доставка завершена"]
                
            # Стилизация и скрытие технической колонки
            styled_df = df_filtered.style.apply(color_rows, axis=1)
            styled_df = styled_df.hide(subset=["Тип подсветки"], axis="columns")
            
            st.dataframe(styled_df, use_container_width=True)
            
            # Генерация и выгрузка Excel
            excel_data = generate_colored_excel(df_output)
            current_time = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            download_name = f"CSE_Результат_{current_time}.xlsx"
            
            st.download_button(
                label="📥 Скачать полный результат с раскраской (Excel)",
                data=excel_data,
                file_name=download_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

if __name__ == "__main__":
    main()
